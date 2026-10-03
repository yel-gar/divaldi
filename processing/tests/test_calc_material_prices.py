"""
Coverage for reading material prices out of the Excel template.

``_load_material_prices`` reads the ``Цены на металл`` sheet in ``data_only``
mode, so it sees cached formula results rather than formulas. That has one
consequence worth pinning: a template that was written but never recalculated by
Excel yields no prices at all, and the module raises rather than silently
producing an offer priced at zero.
"""

import io

import pytest
from openpyxl import Workbook

from processing.calculator.calc import (
    COL_PRICE_NAME,
    COL_PRICE_VALUE,
    PRICE_FIRST_ROW,
    SHEET_PRICES,
    _load_material_prices,
)


def _template_with(rows: list[tuple[object, object]]) -> bytes:
    """Build a minimal workbook whose price sheet holds ``(name, value)`` pairs.

    ``value`` of ``None`` leaves the cell empty, which is what a formula that
    has never been recalculated looks like to ``data_only=True``.
    """
    wb = Workbook()
    wb.active.title = SHEET_PRICES
    sheet = wb[SHEET_PRICES]
    for offset, (name, value) in enumerate(rows):
        row = PRICE_FIRST_ROW + offset
        sheet.cell(row=row, column=COL_PRICE_NAME).value = name
        if value is not None:
            sheet.cell(row=row, column=COL_PRICE_VALUE).value = value

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_load_material_prices_reads_name_value_pairs():
    prices = _load_material_prices(_template_with([("Сталь 3 мм", 1200.0), ("Сталь 5 мм", 2400.5)]))

    assert prices == {"Сталь 3 мм": 1200.0, "Сталь 5 мм": 2400.5}


def test_load_material_prices_strips_whitespace_from_names():
    """Trailing spaces in a hand-edited sheet must not become part of the key."""
    prices = _load_material_prices(_template_with([("  Сталь 3 мм  ", 1200.0)]))

    assert prices == {"Сталь 3 мм": 1200.0}


def test_load_material_prices_converts_numeric_strings():
    """A price stored as text but written in a parseable form is still usable."""
    prices = _load_material_prices(_template_with([("Сталь", "1500.5")]))

    assert prices == {"Сталь": 1500.5}


def test_load_material_prices_rejects_comma_decimal_separators():
    """A Russian-locale Excel writes "1500,50", which Python's float() rejects.

    The row is skipped rather than guessed at, so the material ends up missing
    from the price map and the caller falls back to its own source.
    """
    prices = _load_material_prices(_template_with([("Сталь", "1500,50")]))

    assert prices == {}


def test_load_material_prices_skips_non_numeric_values():
    """An unparseable price is skipped; the surrounding rows still load."""
    prices = _load_material_prices(
        _template_with([("Сталь 3 мм", 1200.0), ("Мусор", "не число"), ("Сталь 5 мм", 2400.0)])
    )

    assert prices == {"Сталь 3 мм": 1200.0, "Сталь 5 мм": 2400.0}


def test_load_material_prices_skips_rows_without_a_name():
    """A value with no material name is not a price entry."""
    prices = _load_material_prices(_template_with([("", 999.0), ("Сталь", 1200.0)]))

    assert prices == {"Сталь": 1200.0}


def test_load_material_prices_raises_when_no_value_was_cached():
    """A template with names but no cached values is a configuration error.

    This is the case where Excel wrote formulas but never recalculated them, so
    `data_only=True` yields None for every price. Producing an offer with zero
    prices would be worse than failing loudly.
    """
    template = _template_with([("Сталь 3 мм", None), ("Сталь 5 мм", None)])

    with pytest.raises(ValueError, match="must be recalculated and saved by Excel"):
        _load_material_prices(template)


def test_load_material_prices_empty_sheet_is_not_an_error():
    """No names at all is a blank template, not a recalculation problem."""
    assert _load_material_prices(_template_with([])) == {}
