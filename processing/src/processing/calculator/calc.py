"""
Excel template filler module for commercial offer generation.

Fills an Excel template (provided as bytes) with position data and returns
the result as bytes. No file I/O is performed — the caller is responsible
for loading the template from storage (e.g. S3) and saving the result.
"""

import io
import math
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook


SHEET_CALC = "Расчёт"
SHEET_PRICES = "Цены на металл"

FIRST_POSITION_ROW = 3
POSITION_BLOCK_HEIGHT = 9

ROW_LASER = 0
ROW_BENDING = 1
ROW_TURNING = 2
ROW_WELDING = 3
ROW_PAINTING = 4

COL_MATERIAL = 3       # C
COL_PRICE_PER_M2 = 4   # D
COL_AREA = 5           # E
COL_HOURS = 8          # H

PRICE_FIRST_ROW = 3
COL_PRICE_NAME = 1     # A
COL_PRICE_VALUE = 8    # H

REQUIRED_FIELDS = (
    "name",
    "material",
    "area_m2",
    "laser_m",
    "bends",
    "welding_m",
    "turning_hours",
    "painting_m2",
)

NUMERIC_FIELDS = (
    "area_m2",
    "laser_m",
    "bends",
    "welding_m",
    "turning_hours",
    "painting_m2",
)



@dataclass(frozen=True)
class Parameters:
    """
    Production norms and rates used when filling the template.

    Values are used to convert physical measurements (meters, bends) into
    working hours, which are then consumed by Excel formulas.
    """

    laser_speed_m_per_hour: float = 10.0
    welding_speed_m_per_hour: float = 2.0
    bending_rate_per_hour: float = 84.0
    painting_rate_m2_per_hour: float = 5.53
    max_positions: int = 10


DEFAULT_PARAMETERS = Parameters()


def _load_material_prices(template_bytes: bytes) -> dict[str, float]:
    """
        Extract material prices from the template's price sheet.

        Reads the file in data_only mode so that formula results (not formulas
        themselves) are returned. This requires the template to have been
        recalculated and saved by Excel at least once.

        Args:
            template_bytes: Raw content of the Excel template.

        Returns:
            Mapping of material name to price per m².
        """
    wb = load_workbook(io.BytesIO(template_bytes), data_only=True)
    sheet = wb[SHEET_PRICES]
    prices: dict[str, float] = {}
    names_seen = 0
    values_missing = 0
    for row in range(PRICE_FIRST_ROW, sheet.max_row + 1):
        name = sheet.cell(row=row, column=COL_PRICE_NAME).value
        value = sheet.cell(row=row, column=COL_PRICE_VALUE).value
        if not name:
            continue
        names_seen += 1
        if value is None:
            values_missing += 1
            continue
        try:
            prices[str(name).strip()] = float(value)
        except (ValueError, TypeError):
            continue

    if names_seen and not prices and values_missing:
        raise ValueError(
            "Material prices sheet has formulas but no cached values. "
            "The template must be recalculated and saved by Excel "
            "before upload, or material_prices must be passed explicitly."
        )
    return prices


def _validate_positions(
    positions: list[dict[str, Any]],
    material_prices: dict[str, float],
    params: Parameters,
) -> None:
    """
    Validate all positions before any work is performed.

    Args:
        positions: List of position dictionaries.
        material_prices: Mapping of known material names to prices.
        params: Production parameters (used for max_positions).

    Raises:
        ValueError: If any position is malformed or references unknown material.
    """
    if len(positions) > params.max_positions:
        raise ValueError(
            f"'positions' contains {len(positions)} items, "
            f"but only {params.max_positions} are supported."
        )

    for idx, pos in enumerate(positions):
        for field in REQUIRED_FIELDS:
            if field not in pos:
                raise ValueError(
                    f"Position {idx + 1} missing required field: {field}"
                )

        material = pos["material"]
        if not isinstance(material, str):
            raise ValueError(
                f"Position {idx + 1}: 'material' must be a string, "
                f"got {type(material).__name__}"
            )
        material = material.strip()
        if material not in material_prices:
            raise ValueError(
                f"Position {idx + 1}: unknown material '{material}'"
            )

        for field in NUMERIC_FIELDS:
            value = pos[field]
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                raise ValueError(
                    f"Position {idx + 1}: field '{field}' must be a real "
                    f"number, got {type(value).__name__}: {value}"
                )
            if not math.isfinite(value) or value < 0:
                raise ValueError(
                    f"Position {idx + 1}: field '{field}' must be a finite "
                    f"number >= 0, got: {value}"
                )


def _write_position(
    sheet,
    block_index: int,
    pos_data: dict[str, Any],
    material_prices: dict[str, float],
    params: Parameters,
) -> None:
    """
    Write a single position's data into its block on the 'Расчёт' sheet.

    Args:
        sheet: Target worksheet.
        block_index: Zero-based index of the position block.
        pos_data: Validated position dictionary.
        material_prices: Mapping of material names to prices per m².
        params: Production parameters used to convert to hours.
    """
    base = FIRST_POSITION_ROW + block_index * POSITION_BLOCK_HEIGHT

    material = pos_data["material"].strip()
    area_m2 = float(pos_data["area_m2"])

    sheet.cell(row=base + ROW_LASER, column=COL_MATERIAL).value = material
    sheet.cell(row=base + ROW_LASER, column=COL_AREA).value = area_m2
    sheet.cell(row=base + ROW_LASER, column=COL_PRICE_PER_M2).value = (
        material_prices[material]
    )

    laser_m = float(pos_data["laser_m"])
    hours_laser = (
        laser_m / params.laser_speed_m_per_hour
        if params.laser_speed_m_per_hour > 0
        else 0.0
    )
    sheet.cell(row=base + ROW_LASER, column=COL_HOURS).value = hours_laser

    bends = float(pos_data["bends"])
    hours_bending = (
        bends / params.bending_rate_per_hour
        if params.bending_rate_per_hour > 0
        else 0.0
    )
    sheet.cell(row=base + ROW_BENDING, column=COL_HOURS).value = hours_bending

    sheet.cell(row=base + ROW_TURNING, column=COL_HOURS).value = float(
        pos_data["turning_hours"]
    )

    welding_m = float(pos_data["welding_m"])
    hours_welding = (
        welding_m / params.welding_speed_m_per_hour
        if params.welding_speed_m_per_hour > 0
        else 0.0
    )
    sheet.cell(row=base + ROW_WELDING, column=COL_HOURS).value = hours_welding

    sheet.cell(row=base + ROW_PAINTING, column=COL_AREA).value = float(
        pos_data["painting_m2"]
    )


def _clear_positions(sheet, params: Parameters) -> None:
    """Reset all input cells for every position block."""
    for i in range(params.max_positions):
        base = FIRST_POSITION_ROW + i * POSITION_BLOCK_HEIGHT
        for col in (COL_MATERIAL, COL_PRICE_PER_M2, COL_AREA, COL_HOURS):
            sheet.cell(row=base + ROW_LASER, column=col).value = None
        for row_offset in (ROW_BENDING, ROW_TURNING, ROW_WELDING):
            sheet.cell(row=base + row_offset, column=COL_HOURS).value = None
        sheet.cell(row=base + ROW_PAINTING, column=COL_AREA).value = None


# api


def process_calculation(
    data: dict[str, Any],
    template_bytes: bytes,
    parameters: Parameters | None = None,
    material_prices: dict[str, float] | None = None,
) -> bytes:
    """
    ...
    Args:
        data: Parsed payload from GigaChat.
        template_bytes: Raw content of the Excel template.
        parameters: Optional production parameters. Falls back to
            ``DEFAULT_PARAMETERS``.
        material_prices: Optional mapping of material name to price per m².
            When omitted, prices are read from the template's
            "Цены на металл" sheet via ``_load_material_prices``.
            Pass explicitly to decouple from the template (e.g. in tests
            or when running openpyxl round-trips that drop formula caches).
    ...
    """
    params = parameters or DEFAULT_PARAMETERS

    if not isinstance(data, dict):
        raise ValueError("'data' must be a dict.")

    positions = data.get("positions")
    if not isinstance(positions, list):
        raise ValueError("'data' must contain a 'positions' list.")
    if not positions:
        raise ValueError("'positions' is empty.")
    if not all(isinstance(p, dict) for p in positions):
        raise ValueError("'positions' must be a list of objects.")

    if material_prices is None:
        material_prices = _load_material_prices(template_bytes)

    _validate_positions(positions, material_prices, params)

    wb = load_workbook(io.BytesIO(template_bytes))
    sheet = wb[SHEET_CALC]
    _clear_positions(sheet, params)

    for idx, pos in enumerate(positions):
        _write_position(sheet, idx, pos, material_prices, params)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
