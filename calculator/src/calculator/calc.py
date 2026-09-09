"""
Excel template filler module for commercial offer generation.

Copies the original Excel template, writes part parameters (material, area,
cutting length, bends, welding, turning hours, painting area) to the "Расчёт"
sheet, and saves the result to the specified output file. Excel formulas
will recalculate automatically when the file is opened by the user.
"""

import json
import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# Constants
LASER_SPEED = 10.0  # meters per hour (laser cutting)
WELDING_SPEED = 2.0  # meters per hour
BENDING_RATE = 84.0  # bends per hour
PAINTING_RATE = 5.53  # m^2 per hour (powder coating)

MAX_POSITIONS = 10


def _load_material_prices(file_path: str) -> dict[str, float]:
    """
    Loads material prices from the 'Цены на металл' sheet of the given Excel file.

    Opens the file with data_only=True to get calculated values (not formulas).

    Args:
        file_path: Path to the Excel file.

    Returns:
        A dictionary {material_name: price_per_m²}. Prices are taken from
        column H ('Стоимость 1м2') starting from row 3.
    """
    wb = load_workbook(file_path, data_only=True)
    sheet = wb["Цены на металл"]
    prices = {}
    for row in range(3, sheet.max_row + 1):
        name_cell = sheet.cell(row=row, column=1)  # A
        price_cell = sheet.cell(row=row, column=8)  # H
        if name_cell.value and price_cell.value:
            try:
                price = float(price_cell.value)
                prices[str(name_cell.value).strip()] = price
            except (ValueError, TypeError):
                continue
    return prices


def _validate_positions(
    positions: list[dict[str, Any]], material_prices: dict[str, float]
) -> None:
    """
    Validates all positions before any file operations.

    Args:
        positions: List of position dictionaries.
        material_prices: Dictionary of valid material names and prices.

    Raises:
        ValueError: If any position is malformed or contains unknown material.
    """
    if len(positions) > MAX_POSITIONS:
        raise ValueError(
            f"JSON 'positions' contains {len(positions)} items, "
            f"but only {MAX_POSITIONS} are supported."
        )

    required_fields = [
        "name",
        "material",
        "area_m2",
        "laser_m",
        "bends",
        "welding_m",
        "turning_hours",
        "painting_m2",
    ]
    numeric_fields = [
        "area_m2",
        "laser_m",
        "bends",
        "welding_m",
        "turning_hours",
        "painting_m2",
    ]

    for idx, pos in enumerate(positions):
        for field in required_fields:
            if field not in pos:
                raise ValueError(f"Position {idx+1} missing required field: {field}")

        material = pos.get("material", "").strip()
        if material not in material_prices:
            raise ValueError(
                f"Position {idx+1}: unknown material '{material}'. "
                f"Available materials: {', '.join(list(material_prices.keys())[:5])}..."
            )

        for field in numeric_fields:
            try:
                float(pos[field])
            except (ValueError, TypeError):
                raise ValueError(
                    f"Position {idx+1}: field '{field}' is not a number: {pos[field]}"
                )


def _write_position(
    sheet, row_offset: int, pos_data: dict[str, Any], material_prices: dict[str, float]
) -> None:
    """
    Writes a single part position data to the 'Расчёт' sheet.
    """
    row_laser = row_offset  # laser cutting
    row_bending = row_offset + 1  # bending
    row_turning = row_offset + 2  # turning
    row_welding = row_offset + 3  # welding
    row_painting = row_offset + 4  # powder coating

    material = pos_data.get("material", "").strip()
    area = float(pos_data.get("area_m2", 0.0))

    sheet.cell(row=row_laser, column=3).value = material  # C
    sheet.cell(row=row_laser, column=5).value = area  # E

    price_per_m2 = material_prices[material]
    sheet.cell(row=row_laser, column=4).value = price_per_m2  # D

    laser_m = float(pos_data.get("laser_m", 0.0))
    hours_laser = laser_m / LASER_SPEED if LASER_SPEED > 0 else 0.0
    sheet.cell(row=row_laser, column=8).value = hours_laser

    bends = float(pos_data.get("bends", 0.0))
    hours_bending = bends / BENDING_RATE if BENDING_RATE > 0 else 0.0
    sheet.cell(row=row_bending, column=8).value = hours_bending

    hours_turning = float(pos_data.get("turning_hours", 0.0))
    sheet.cell(row=row_turning, column=8).value = hours_turning

    welding_m = float(pos_data.get("welding_m", 0.0))
    hours_welding = welding_m / WELDING_SPEED if WELDING_SPEED > 0 else 0.0
    sheet.cell(row=row_welding, column=8).value = hours_welding

    painting_m2 = float(pos_data.get("painting_m2", 0.0))
    sheet.cell(row=row_painting, column=5).value = painting_m2  # E


def process_calculation(json_data: str, template_path: str, output_path: str) -> str:
    """
    Processes JSON data from GigaChat and creates a filled Excel commercial offer.

    Workflow:
        1. Parses and validates the JSON.
        2. Loads material prices from the template (read-only, data_only=True).
        3. Validates all positions against the material prices and numeric fields.
        4. Copies the template to the output path.
        5. Opens the copy for writing, clears old data, writes new positions.
        6. Saves the file.

    Args:
        json_data: JSON string from GigaChat.
            Expected structure:
            {
                "positions": [
                    {
                        "name": "Bracket left",
                        "material": "1,5 мм Оц. Сталь (2500х1250)",
                        "area_m2": 0.8,
                        "laser_m": 2.5,
                        "bends": 12,
                        "welding_m": 0.5,
                        "turning_hours": 0.0,
                        "painting_m2": 0.8
                    }
                ]
            }
        template_path: Path to the Excel template file (original).
        output_path: Path where the result file will be saved (including filename).

    Returns:
        The path to the created Excel file (output_path).

    Raises:
        ValueError: If JSON is invalid, structure is wrong, positions exceed limit,
                    or any position contains unknown material or non-numeric values.
        FileNotFoundError: If the template file does not exist.
        PermissionError: If write permissions are insufficient.
    """
    try:
        data = json.loads(json_data)
    except json.JSONDecodeError as e:
        raise ValueError(f"Invalid JSON: {e}")

    if not isinstance(data, dict):
        raise ValueError("JSON top-level value must be an object.")

    positions = data.get("positions", [])
    if not isinstance(positions, list) or not all(
        isinstance(pos, dict) for pos in positions
    ):
        raise ValueError("JSON 'positions' must be a list of objects.")
    if not positions:
        raise ValueError("JSON missing 'positions' key or it is empty.")

    template_file = Path(template_path)
    if not template_file.exists():
        raise FileNotFoundError(f"Template file not found: {template_path}")

    material_prices = _load_material_prices(str(template_file))

    _validate_positions(positions, material_prices)

    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_file, output_file)

    wb = load_workbook(output_file)
    calc_sheet = wb["Расчёт"]

    for i in range(10):
        row_start = 3 + i * 9
        # Laser: columns C, D, E, H
        for col in [3, 4, 5, 8]:
            calc_sheet.cell(row=row_start, column=col).value = None
        # Bending: H
        calc_sheet.cell(row=row_start + 1, column=8).value = None
        # Turning: H
        calc_sheet.cell(row=row_start + 2, column=8).value = None
        # Welding: H
        calc_sheet.cell(row=row_start + 3, column=8).value = None
        # Painting: E
        calc_sheet.cell(row=row_start + 4, column=5).value = None

    for idx, pos in enumerate(positions):
        row_start = 3 + idx * 9
        _write_position(calc_sheet, row_start, pos, material_prices)

    wb.save(output_file)

    return str(output_file)
