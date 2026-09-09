"""
Tests for calc module (Excel template filler).
"""

import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

# Add src to Python path for imports
import sys
sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from calculator.calc import (
    process_calculation,
    _load_material_prices,
    _write_position,
    LASER_SPEED,
    WELDING_SPEED,
    BENDING_RATE,
    PAINTING_RATE,
)


# ---- Fixtures ----

@pytest.fixture
def template_path(tmp_path) -> str:
    """
    Copy the real template file (calc.xlsx) to a temporary location
    and return its path.
    """
    # The template is located in src/calculator/excel_calc/ relative to project root
    project_root = Path(__file__).parent.parent
    possible_paths = [
        project_root / "src" / "calculator" / "excel_calc" / "calc.xlsx",
        project_root / "excel_calc" / "calc.xlsx",
        project_root.parent / "excel_calc" / "calc.xlsx",
    ]
    original = None
    for p in possible_paths:
        if p.exists():
            original = p
            break
    if original is None:
        pytest.skip("Template file calc.xlsx not found in any expected location")

    temp_template = tmp_path / "calc.xlsx"
    shutil.copy2(original, temp_template)
    return str(temp_template)


@pytest.fixture
def sample_json_single() -> str:
    """Valid JSON with a single position."""
    return json.dumps({
        "positions": [
            {
                "name": "Bracket support",
                "material": "1,5 мм Оц. Сталь (2500х1250)",
                "area_m2": 0.8,
                "laser_m": 2.5,
                "bends": 12,
                "welding_m": 0.5,
                "turning_hours": 0.0,
                "painting_m2": 0.8
            }
        ]
    })


@pytest.fixture
def sample_json_multiple() -> str:
    """Valid JSON with 3 positions."""
    return json.dumps({
        "positions": [
            {
                "name": "Part A",
                "material": "1,5 мм Оц. Сталь (2500х1250)",
                "area_m2": 0.8,
                "laser_m": 2.5,
                "bends": 12,
                "welding_m": 0.5,
                "turning_hours": 0.0,
                "painting_m2": 0.8
            },
            {
                "name": "Part B",
                "material": "2,0 мм Оц. Сталь (2500х1250)",
                "area_m2": 1.2,
                "laser_m": 3.0,
                "bends": 8,
                "welding_m": 0.0,
                "turning_hours": 0.5,
                "painting_m2": 0.0
            },
            {
                "name": "Part C",
                "material": "3,0 мм Ст3 (3000х1500)",
                "area_m2": 2.0,
                "laser_m": 4.0,
                "bends": 0,
                "welding_m": 1.0,
                "turning_hours": 0.0,
                "painting_m2": 1.5
            }
        ]
    })


@pytest.fixture
def output_path(tmp_path) -> str:
    """Return a path for the output file in the temporary directory."""
    return str(tmp_path / "result.xlsx")


# ---- Tests for process_calculation ----

def test_process_calculation_success(
    template_path: str,
    sample_json_single: str,
    output_path: str
) -> None:
    """Test successful processing of a single position."""
    result_path = process_calculation(sample_json_single, template_path, output_path)
    assert result_path == output_path
    assert Path(output_path).exists()

    wb = load_workbook(output_path)
    calc_sheet = wb["Расчёт"]

    assert calc_sheet.cell(row=3, column=3).value == "1,5 мм Оц. Сталь (2500х1250)"
    assert calc_sheet.cell(row=3, column=5).value == 0.8
    # Price per m² should be a number (not None)
    price = calc_sheet.cell(row=3, column=4).value
    assert price is not None and price > 0
    assert calc_sheet.cell(row=3, column=8).value == pytest.approx(0.25)
    assert calc_sheet.cell(row=4, column=8).value == pytest.approx(12 / 84)
    assert calc_sheet.cell(row=5, column=8).value == 0.0
    assert calc_sheet.cell(row=6, column=8).value == pytest.approx(0.25)
    assert calc_sheet.cell(row=7, column=5).value == 0.8
    assert "КП с ндс" in wb.sheetnames


def test_process_calculation_multiple_positions(
    template_path: str,
    sample_json_multiple: str,
    output_path: str
) -> None:
    """Test processing multiple positions (up to 10)."""
    process_calculation(sample_json_multiple, template_path, output_path)
    assert Path(output_path).exists()

    wb = load_workbook(output_path)
    calc_sheet = wb["Расчёт"]

    assert calc_sheet.cell(row=3, column=3).value == "1,5 мм Оц. Сталь (2500х1250)"
    assert calc_sheet.cell(row=12, column=3).value == "2,0 мм Оц. Сталь (2500х1250)"
    assert calc_sheet.cell(row=21, column=3).value == "3,0 мм Ст3 (3000х1500)"
    assert calc_sheet.cell(row=30, column=3).value is None
    assert calc_sheet.cell(row=30, column=5).value is None


def test_process_calculation_clears_old_data(
    template_path: str,
    sample_json_single: str,
    output_path: str
) -> None:
    """Test that old data is cleared before writing new data."""
    process_calculation(sample_json_single, template_path, output_path)

    new_json = json.dumps({
        "positions": [
            {
                "name": "Another part",
                "material": "2,0 мм Оц. Сталь (2500х1250)",
                "area_m2": 1.0,
                "laser_m": 1.0,
                "bends": 0,
                "welding_m": 0.0,
                "turning_hours": 0.0,
                "painting_m2": 0.0
            }
        ]
    })
    process_calculation(new_json, template_path, output_path)

    wb = load_workbook(output_path)
    calc_sheet = wb["Расчёт"]

    assert calc_sheet.cell(row=3, column=3).value == "2,0 мм Оц. Сталь (2500х1250)"
    assert calc_sheet.cell(row=3, column=5).value == 1.0
    assert calc_sheet.cell(row=3, column=8).value == pytest.approx(1.0 / LASER_SPEED)
    # Second position should be empty
    assert calc_sheet.cell(row=12, column=3).value is None
    assert calc_sheet.cell(row=12, column=5).value is None
    assert calc_sheet.cell(row=12, column=8).value is None


def test_process_calculation_creates_output_directory(
    template_path: str,
    sample_json_single: str,
    tmp_path: Path
) -> None:
    """Test that the output directory is created if it doesn't exist."""
    deep_path = tmp_path / "sub" / "dir" / "result.xlsx"
    output_path = str(deep_path)
    assert not deep_path.parent.exists()
    process_calculation(sample_json_single, template_path, output_path)
    assert deep_path.exists()


# ---- Error handling tests ----

def test_process_calculation_invalid_json(template_path: str, output_path: str) -> None:
    with pytest.raises(ValueError, match="Invalid JSON"):
        process_calculation("{not json", template_path, output_path)


def test_process_calculation_missing_positions(
    template_path: str,
    output_path: str
) -> None:
    json_no_positions = json.dumps({"other": "data"})
    with pytest.raises(ValueError, match="missing 'positions' key"):
        process_calculation(json_no_positions, template_path, output_path)


def test_process_calculation_empty_positions(
    template_path: str,
    output_path: str
) -> None:
    json_empty = json.dumps({"positions": []})
    with pytest.raises(ValueError, match="missing 'positions' key or it is empty"):
        process_calculation(json_empty, template_path, output_path)


def test_process_calculation_template_not_found(
    sample_json_single: str,
    output_path: str
) -> None:
    with pytest.raises(FileNotFoundError, match="Template file not found"):
        process_calculation(sample_json_single, "/non/existent/file.xlsx", output_path)


def test_process_calculation_missing_material(
    template_path: str,
    output_path: str
) -> None:
    json_with_unknown = json.dumps({
        "positions": [
            {
                "name": "Unknown",
                "material": "Non-existent material",
                "area_m2": 1.0,
                "laser_m": 1.0,
                "bends": 0,
                "welding_m": 0.0,
                "turning_hours": 0.0,
                "painting_m2": 0.0
            }
        ]
    })
    result = process_calculation(json_with_unknown, template_path, output_path)
    wb = load_workbook(result)
    calc_sheet = wb["Расчёт"]
    # Price per m² should be 0 (fallback)
    assert calc_sheet.cell(row=3, column=4).value == 0.0


# ---- Tests for internal functions ----

def test_load_material_prices(template_path: str) -> None:
    """Test that _load_material_prices reads calculated values correctly."""
    prices = _load_material_prices(template_path)
    assert isinstance(prices, dict)
    # Check that at least one known material exists and has a positive price
    known_materials = [
        "1,5 мм Оц. Сталь (2500х1250)",
        "2,0 мм Оц. Сталь (2500х1250)",
        "3,0 мм Ст3 (3000х1500)"
    ]
    found_any = False
    for mat in known_materials:
        if mat in prices and prices[mat] > 0:
            found_any = True
            break
    assert found_any, f"No known materials found in {list(prices.keys())[:5]}"


def test_write_position(template_path: str) -> None:
    """Test the _write_position function directly."""
    wb = load_workbook(template_path)
    calc_sheet = wb["Расчёт"]
    material_prices = _load_material_prices(template_path)  # use path to get prices

    # Ensure we have at least one price to test with
    assert material_prices, "No material prices loaded"

    pos_data = {
        "material": "1,5 мм Оц. Сталь (2500х1250)",
        "area_m2": 0.8,
        "laser_m": 2.5,
        "bends": 12,
        "welding_m": 0.5,
        "turning_hours": 0.0,
        "painting_m2": 0.8
    }

    _write_position(calc_sheet, 3, pos_data, material_prices)

    assert calc_sheet.cell(row=3, column=3).value == "1,5 мм Оц. Сталь (2500х1250)"
    assert calc_sheet.cell(row=3, column=5).value == 0.8
    expected_price = material_prices.get("1,5 мм Оц. Сталь (2500х1250)", 0.0)
    assert calc_sheet.cell(row=3, column=4).value == expected_price
    assert calc_sheet.cell(row=3, column=8).value == pytest.approx(2.5 / LASER_SPEED)
    assert calc_sheet.cell(row=4, column=8).value == pytest.approx(12 / BENDING_RATE)
    assert calc_sheet.cell(row=5, column=8).value == 0.0
    assert calc_sheet.cell(row=6, column=8).value == pytest.approx(0.5 / WELDING_SPEED)
    assert calc_sheet.cell(row=7, column=5).value == 0.8


# ---- Integration test ----

def test_integration_with_real_template(
    template_path: str,
    sample_json_single: str,
    tmp_path: Path
) -> None:
    out_file = tmp_path / "integration_result.xlsx"
    process_calculation(sample_json_single, template_path, str(out_file))
    assert out_file.exists()
    wb = load_workbook(out_file)
    kp_sheet = wb["КП с ндс"]
    assert kp_sheet.cell(row=15, column=2).value == "=Расчёт!C3"
    assert kp_sheet.cell(row=15, column=4).value == "=Расчёт!L8"
